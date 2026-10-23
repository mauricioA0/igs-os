"""Write the import templates (one CSV per sheet + one workbook) from the specs.

uv run igs-import templates docs/import/templates
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

from igs_common.importer.schema import APPLY_ORDER, SPECS, Entity, FieldType


def template_rows(entity: Entity) -> tuple[list[str], list[str]]:
    spec = SPECS[entity]
    return [f.name for f in spec.fields], [f.example for f in spec.fields]


def csv_template(entity: Entity) -> str:
    headers, example = template_rows(entity)
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(headers)
    writer.writerow(example)
    return buf.getvalue()


def xlsx_template() -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font

    wb = Workbook()
    readme = wb.active
    assert readme is not None
    readme.title = "README"
    readme.append(["IGS OS import template"])
    readme.append(["One sheet per kind of data. Keep the header row; delete the example row."])
    readme.append(["Bold headers are required. Empty cells never erase what IGS OS has."])
    readme.append(["Dates: YYYY-MM-DD. Amounts: whole US dollars."])
    readme.append([])
    readme.append(["sheet", "what it holds", "rows match by"])
    for entity in APPLY_ORDER:
        spec = SPECS[entity]
        readme.append([entity.value, spec.description, " + ".join(spec.key)])
    readme["A1"].font = Font(bold=True, size=14)
    for entity in APPLY_ORDER:
        spec = SPECS[entity]
        ws = wb.create_sheet(entity.value)
        headers, example = template_rows(entity)
        ws.append(headers)
        ws.append(example)
        for col, f in enumerate(spec.fields, start=1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=f.required)
            note = f.description
            if f.type is FieldType.ENUM:
                note += "\nOne of: " + ", ".join(f.choices)
            cell.comment = Comment(note, "IGS OS")
            ws.column_dimensions[cell.column_letter].width = max(14, len(f.name) + 4)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def write_templates(directory: Path) -> list[Path]:
    directory.mkdir(parents=True, exist_ok=True)
    written = []
    for entity in APPLY_ORDER:
        path = directory / f"{entity.value}.csv"
        path.write_text(csv_template(entity), encoding="utf-8")
        written.append(path)
    path = directory / "igs-os-import.xlsx"
    path.write_bytes(xlsx_template())
    written.append(path)
    return written
