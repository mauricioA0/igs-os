"""Read CSV and XLSX files into plain sheets (header row + string/date cells).

Nothing is interpreted here: cells stay as the spreadsheet had them; `validate` coerces them.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from igs_common.importer.schema import SPECS, Entity

Cell = str | int | float | bool | date | datetime | None

MAX_ROWS = 20_000


class ImportFileError(ValueError):
    """The file can't be read at all (wrong format, no header, too big)."""


@dataclass
class Sheet:
    name: str
    headers: list[str]
    # Spreadsheet row number (1-based, header = 1) → cells aligned with `headers`.
    rows: list[tuple[int, list[Cell]]] = field(default_factory=list)
    entity: Entity | None = None


def normalize(header: str) -> str:
    """'Fecha de Inicio ' → 'fecha_de_inicio'; accents and punctuation dropped."""
    text = unicodedata.normalize("NFKD", str(header)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def detect_entity(name: str) -> Entity | None:
    """Sheet or file name → entity ('Clientes 2026.csv' → accounts)."""
    norm = normalize(Path(name).stem if "." in name else name)
    for entity, spec in SPECS.items():
        for alias in (entity.value, *spec.sheet_aliases):
            if norm == alias or norm.startswith(alias + "_") or norm.endswith("_" + alias):
                return entity
    return None


def read_file(data: bytes, filename: str, *, entity: Entity | None = None) -> list[Sheet]:
    suffix = Path(filename).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        sheets = _read_xlsx(data)
    elif suffix in (".csv", ".tsv", ".txt"):
        sheets = [_read_csv(data, Path(filename).stem)]
    else:
        raise ImportFileError(f"unsupported file type {suffix or '(none)'}; use .xlsx or .csv")
    for sheet in sheets:
        sheet.entity = entity if entity is not None and len(sheets) == 1 else None
        if sheet.entity is None:
            sheet.entity = detect_entity(sheet.name)
    return sheets


def _read_csv(data: bytes, name: str) -> Sheet:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")  # Excel on Windows exports CSV as cp1252/latin-1
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
            text[:4096], delimiters=",;\t"
        )
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = [r for r in reader]
    return _sheet(name, rows)


def _read_xlsx(data: bytes) -> list[Sheet]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a zoo of zipfile/xml errors
        raise ImportFileError(f"not a readable .xlsx file: {exc}") from None
    try:
        sheets = []
        for ws in wb.worksheets:
            rows: list[list[Cell]] = [[_cell(c) for c in r] for r in ws.iter_rows(values_only=True)]
            if any(any(c not in (None, "") for c in r) for r in rows):
                sheets.append(_sheet(ws.title, rows))
        return sheets
    finally:
        wb.close()


def _cell(value: object) -> Cell:
    """openpyxl hands back str/number/bool/date/datetime/time; anything else becomes text."""
    if value is None or isinstance(value, (str, int, float, bool, date, datetime)):
        return value
    return str(value)


def _sheet(name: str, rows: Sequence[Sequence[Cell]]) -> Sheet:
    # Skip blank/title rows above the header: the header is the first row with ≥2 cells.
    start = next(
        (i for i, r in enumerate(rows) if sum(1 for c in r if c not in (None, "")) >= 2), None
    )
    if start is None:
        raise ImportFileError(f"sheet {name!r} has no header row")
    headers = [str(c).strip() if c is not None else "" for c in rows[start]]
    body: list[tuple[int, list[Cell]]] = []
    for offset, raw in enumerate(rows[start + 1 :], start=start + 2):
        cells: list[Cell] = [
            (c.strip() or None) if isinstance(c, str) else c for c in list(raw)[: len(headers)]
        ]
        cells += [None] * (len(headers) - len(cells))
        if all(c is None for c in cells):
            continue
        body.append((offset, cells))
        if len(body) > MAX_ROWS:
            raise ImportFileError(f"sheet {name!r} has more than {MAX_ROWS} rows")
    return Sheet(name=name, headers=headers, rows=body)
