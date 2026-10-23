"""Turn source bytes (PDF, DOCX, HTML, text/markdown) or an interaction into a ParsedDocument.

Parsers keep structure (heading paths) so chunks never straddle sections and citations can say
"Cancellation policy › Refunds". They never execute or follow anything inside the document.
"""

import io
import re
from collections.abc import Mapping
from datetime import datetime
from html.parser import HTMLParser
from typing import Any

from .types import ParsedDocument, Section


class UnsupportedFormatError(ValueError):
    """Raised for MIME types we don't parse (non-retryable for the ingestion workflow)."""


PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
HTML = "text/html"
TEXT = "text/plain"
MARKDOWN = "text/markdown"

SUPPORTED_MIME_TYPES = frozenset({PDF, DOCX, HTML, TEXT, MARKDOWN})


def parse_bytes(data: bytes, mime_type: str, *, title: str) -> ParsedDocument:
    mime = mime_type.split(";", 1)[0].strip().lower()
    if mime == PDF:
        return parse_pdf(data, title=title)
    if mime == DOCX:
        return parse_docx(data, title=title)
    if mime in {HTML, "application/xhtml+xml"}:
        return parse_html(_decode(data), title=title)
    if mime in {TEXT, MARKDOWN, "text/x-markdown"}:
        return parse_markdown(_decode(data), title=title)
    raise UnsupportedFormatError(f"unsupported mime type: {mime_type}")


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")  # pragma: no cover - latin-1 never fails


# ── Markdown / plain text ─────────────────────────────────────────────────

_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")


def parse_markdown(text: str, *, title: str) -> ParsedDocument:
    """ATX headings (``#``…``######``) define sections; plain text is one untitled section."""
    sections: list[Section] = []
    path: list[tuple[int, str]] = []
    buf: list[str] = []

    def flush() -> None:
        body = "\n".join(buf).strip()
        if body:
            sections.append(Section(tuple(h for _, h in path), body))
        buf.clear()

    for line in text.splitlines():
        m = _MD_HEADING.match(line)
        if m:
            flush()
            level = len(m.group(1))
            while path and path[-1][0] >= level:
                path.pop()
            path.append((level, m.group(2).strip()))
        else:
            buf.append(line)
    flush()
    return ParsedDocument(title=title, sections=tuple(sections))


# ── HTML ──────────────────────────────────────────────────────────────────

_BLOCK_TAGS = {"p", "li", "div", "section", "article", "td", "th", "blockquote", "pre", "br", "tr"}
_SKIP_TAGS = {"script", "style", "noscript", "template", "svg", "head", "iframe", "object"}


class _HTMLSections(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.sections: list[Section] = []
        self.path: list[tuple[int, str]] = []
        self.buf: list[str] = []
        self.heading: list[str] | None = None
        self.heading_level = 0
        self.skip_depth = 0
        self.hidden_depth = 0
        self.saw_hidden_text = False
        self._stack: list[tuple[str, bool]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in _SKIP_TAGS:
            self.skip_depth += 1
            return
        if tag in {"br", "img", "hr", "meta", "link", "input"}:
            if tag == "br":
                self.buf.append("\n")
            return
        style = (dict(attrs).get("style") or "").replace(" ", "").lower()
        hidden = (
            "display:none" in style
            or "visibility:hidden" in style
            or any(k == "hidden" for k, _ in attrs)
            or dict(attrs).get("aria-hidden") == "true"
        )
        # Hidden text is kept, not dropped: it is a classic carrier for injected instructions,
        # so it must reach the quarantine check. The document is flagged as containing it.
        self._stack.append((tag, hidden))
        if hidden:
            self.hidden_depth += 1
        if re.fullmatch(r"h[1-6]", tag):
            self._flush()
            self.heading = []
            self.heading_level = int(tag[1])
        elif tag in _BLOCK_TAGS:
            self.buf.append("\n\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in _SKIP_TAGS:
            self.skip_depth = max(0, self.skip_depth - 1)
            return
        while self._stack:
            open_tag, hidden = self._stack.pop()
            if hidden:
                self.hidden_depth -= 1
            if open_tag == tag:
                break
        if re.fullmatch(r"h[1-6]", tag) and self.heading is not None:
            text = " ".join("".join(self.heading).split())
            while self.path and self.path[-1][0] >= self.heading_level:
                self.path.pop()
            if text:
                self.path.append((self.heading_level, text))
            self.heading = None
        elif tag in _BLOCK_TAGS:
            self.buf.append("\n\n")

    def handle_data(self, data: str) -> None:
        if self.skip_depth:
            return
        if self.hidden_depth and data.strip():
            self.saw_hidden_text = True
        if self.heading is not None:
            self.heading.append(data)
        else:
            self.buf.append(data)

    def _flush(self) -> None:
        raw = "".join(self.buf)
        paragraphs = [" ".join(p.split()) for p in re.split(r"\n\s*\n", raw)]
        body = "\n\n".join(p for p in paragraphs if p)
        if body:
            self.sections.append(Section(tuple(h for _, h in self.path), body))
        self.buf.clear()

    def finish(self) -> list[Section]:
        self._flush()
        return self.sections


def parse_html(html: str, *, title: str) -> ParsedDocument:
    parser = _HTMLSections()
    parser.feed(html)
    parser.close()
    m = re.search(r"<title[^>]*>(.*?)</title>", html, re.IGNORECASE | re.DOTALL)
    doc_title = " ".join(m.group(1).split()) if m and not title else title
    sections = tuple(parser.finish())
    return ParsedDocument(
        title=doc_title, sections=sections, metadata={"hidden_text": parser.saw_hidden_text}
    )


# ── DOCX ──────────────────────────────────────────────────────────────────


def parse_docx(data: bytes, *, title: str) -> ParsedDocument:
    from docx import Document as DocxDocument  # local import: optional heavy dependency

    document = DocxDocument(io.BytesIO(data))
    sections: list[Section] = []
    path: list[tuple[int, str]] = []
    buf: list[str] = []

    def flush() -> None:
        body = "\n\n".join(p for p in buf if p).strip()
        if body:
            sections.append(Section(tuple(h for _, h in path), body))
        buf.clear()

    for para in document.paragraphs:
        text = para.text.strip()
        style = (para.style.name if para.style is not None else "") or ""
        level = _docx_heading_level(style)
        if level and text:
            flush()
            while path and path[-1][0] >= level:
                path.pop()
            path.append((level, text))
        elif text:
            buf.append(text)
    for table in document.tables:
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        buf.append("\n".join(r for r in rows if r.strip(" |")))
    flush()
    core_title = (document.core_properties.title or "").strip()
    return ParsedDocument(title=title or core_title, sections=tuple(sections))


def _docx_heading_level(style: str) -> int:
    if style == "Title":
        return 1
    m = re.match(r"Heading\s*(\d)", style)
    return int(m.group(1)) if m else 0


# ── PDF ───────────────────────────────────────────────────────────────────


def parse_pdf(data: bytes, *, title: str) -> ParsedDocument:
    """One section per page (``("Page 3",)``): PDF has no reliable heading structure, and page
    numbers are what people cite. Extraction failures on a page are skipped, not fatal."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    sections: list[Section] = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001, S112 - a broken page must not fail the document
            continue
        text = _clean_pdf_text(text)
        if text:
            sections.append(Section((f"Page {number}",), text))
    meta_title = ""
    if reader.metadata is not None and reader.metadata.title:
        meta_title = str(reader.metadata.title).strip()
    return ParsedDocument(
        title=title or meta_title, sections=tuple(sections), metadata={"pages": len(reader.pages)}
    )


def _clean_pdf_text(text: str) -> str:
    text = re.sub(r"-\n(?=[a-záéíóúñ])", "", text)  # re-join hyphenated line breaks
    lines = [ln.strip() for ln in text.splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


# ── Interactions (emails, calls, chats, notes) ────────────────────────────


def parse_interaction(
    *,
    channel: str,
    direction: str,
    occurred_at: datetime,
    subject: str | None,
    body: str,
    account_name: str | None = None,
    stakeholder: str | None = None,
    metadata: Mapping[str, Any] | None = None,
) -> ParsedDocument:
    """One section; the heading carries channel/date/who so a retrieved chunk is self-describing
    ("Email · inbound · 2026-09-14 · Robert Haines") for citations and the answer model."""
    who = stakeholder or account_name or ""
    heading = " · ".join(
        p for p in (channel.capitalize(), direction, occurred_at.date().isoformat(), who) if p
    )
    path = (heading, subject.strip()) if subject and subject.strip() else (heading,)
    title = subject.strip() if subject and subject.strip() else heading
    return ParsedDocument(
        title=title,
        sections=(Section(path, body.strip()),) if body.strip() else (),
        metadata=dict(metadata or {}),
    )
